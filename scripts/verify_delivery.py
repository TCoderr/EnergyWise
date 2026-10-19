"""Verify the EnergyWise delivery. Writes reports/delivery_verification.txt (with --zip: reports/zip_verification.txt).
Run: .venv/Scripts/python scripts/verify_delivery.py [--zip]
"""
import hashlib
import json
import re
import sys
import time
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import nbformat
from docx import Document
from pptx import Presentation
from pypdf import PdfReader
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
AUTHOR = 'Tugay Sinan Ergün'
ZIP = ROOT / 'dist' / 'EnergyWise_Teslim.zip'
results = []


def check(name, ok, detail):
    results.append((name, bool(ok), detail))


def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def tr(x, d): return f'{x:.{d}f}'.replace('.', ',')


REQUIRED = ['app.py', 'src/energywise_core.py', 'src/train_model.py', 'notebooks/EnergyWise.ipynb', 'data/energy_efficiency.csv',
            'data/raw/energy_efficiency.zip', 'docs/VERI_SOZLUGU.md', 'models/energywise_linear_regression.joblib',
            'models/model_metadata.json', 'requirements.txt', 'requirements-dev.txt', 'README.md',
            'reports/EnergyWise_Proje_Raporu.pdf', 'reports/EnergyWise_Proje_Raporu.docx', 'presentation/EnergyWise_Sunum.pptx',
            'presentation/EnergyWise_Sunum.pdf', 'reports/console_screenshot.png', 'reports/console_transcript.txt',
            'reports/TEST_REPORT.md', 'reports/test_results.xml', 'tests/test_energywise.py',
            'DELIVERY_CHECKLIST.md', 'FINAL_REPORT.md']
missing = [f for f in REQUIRED if not (ROOT / f).is_file() or (ROOT / f).stat().st_size == 0]
check('Required files', not missing, f'{len(REQUIRED) - len(missing)}/{len(REQUIRED)} present' + (f'; missing {missing}' if missing else ''))

nbm = json.loads((ROOT / 'reports/notebook_metrics.json').read_text(encoding='utf-8'))
meta = json.loads((ROOT / 'models/model_metadata.json').read_text(encoding='utf-8'))

# ---------------------------------------------------------------- notebook
nb = nbformat.read(ROOT / 'notebooks/EnergyWise.ipynb', as_version=4)
code = [c for c in nb.cells if c.cell_type == 'code']
errors = [o for c in code for o in c.get('outputs', []) if o.output_type == 'error']
not_run = [i for i, c in enumerate(code) if c.execution_count is None]
src = '\n'.join(c.source for c in code)
needed = ['.head()', '.shape', '.describe()', '.groupby(', '.value_counts()', '.isnull()', 'train_test_split', 'random_state=42',
          'LinearRegression', 'mean_absolute_error', 'mean_squared_error', 'r2_score', 'model.score', 'model.predict', 'plt.']
miss_calls = [n for n in needed if n not in src]
win_paths = re.findall(r'(?<![A-Za-z])[A-Za-z]:(?:\\\\|/(?!/))', src)   # D:\x or C:/x, but not https://
check('Notebook executed without errors', code and not errors and not not_run, f'{len(code)} code cells, all executed, errors={len(errors)}')
check('Notebook uses required operations', not miss_calls, 'head/shape/describe/groupby/value_counts/isnull/split(42)/LinearRegression/MAE/RMSE/R2/score/predict/plots' + (f'; MISSING {miss_calls}' if miss_calls else ''))
check('Notebook has no Windows-specific paths (Colab)', not win_paths, 'no drive-letter paths in code cells' if not win_paths else f'found {win_paths}')
figs = sorted((ROOT / 'reports/figures').glob('*.png'))
check('At least 2 labelled figures', len(figs) >= 2 and src.count('set_title') >= 2 and src.count('set_xlabel') >= 2,
      f'{len(figs)} PNG figures; titles={src.count("set_title")}, x-labels={src.count("set_xlabel")}')

# ---------------------------------------------------------------- numbers consistent everywhere
same = all(abs(nbm[p][k] - meta[p][k]) < 1e-12 for p in ('train_metrics', 'test_metrics') for k in ('MAE', 'RMSE', 'R2'))
check('Notebook metrics == trained model metadata', same and nbm['n_train'] == meta['n_train'] and nbm['n_test'] == meta['n_test'],
      f"test MAE {nbm['test_metrics']['MAE']:.4f}, RMSE {nbm['test_metrics']['RMSE']:.4f}, R2 {nbm['test_metrics']['R2']:.4f}; split {nbm['n_train']}/{nbm['n_test']}")
TE, TRm = nbm['test_metrics'], nbm['train_metrics']
expect = [tr(TE['MAE'], 3), tr(TE['RMSE'], 3), tr(TE['R2'], 4), tr(TRm['MAE'], 3), tr(TRm['RMSE'], 3), tr(TRm['R2'], 4)] + [tr(v, 2) for v in nbm['new_predictions']]
pdf = PdfReader(ROOT / 'reports/EnergyWise_Proje_Raporu.pdf'); pdf_text = ''.join(p.extract_text() for p in pdf.pages)
miss_pdf = [e for e in expect if e not in pdf_text]
check('Report PDF: 1 page, numbers match notebook', len(pdf.pages) == 1 and not miss_pdf, f'pages={len(pdf.pages)}; {len(expect) - len(miss_pdf)}/{len(expect)} values found' + (f'; MISSING {miss_pdf}' if miss_pdf else ''))
prs = Presentation(ROOT / 'presentation/EnergyWise_Sunum.pptx')
deck_text = ' '.join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
miss_deck = [e for e in expect if e not in deck_text]
deck_pdf = PdfReader(ROOT / 'presentation/EnergyWise_Sunum.pdf')
pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
check('Deck: 3-5 slides, numbers match, PDF exported', 3 <= len(prs.slides) <= 5 and len(deck_pdf.pages) == len(prs.slides) and not miss_deck and pics >= 5,
      f'slides={len(prs.slides)}, pdf pages={len(deck_pdf.pages)}, pictures={pics}' + (f'; MISSING {miss_deck}' if miss_deck else ''))
readme = (ROOT / 'README.md').read_text(encoding='utf-8')
miss_readme = [e for e in expect[:6] if e not in readme]
check('README metrics match', not miss_readme, 'all 6 train/test metrics present' if not miss_readme else f'MISSING {miss_readme}')

# ---------------------------------------------------------------- authorship
doc = Document(ROOT / 'reports/EnergyWise_Proje_Raporu.docx')
authors = [doc.core_properties.author, doc.core_properties.last_modified_by, prs.core_properties.author, (pdf.metadata or {}).get('/Author'),
           (deck_pdf.metadata or {}).get('/Author')]
nb_auth = nb.metadata.get('authors', [{}])[0].get('name')
check('Author metadata = Tugay Sinan Ergün only', all(a == AUTHOR for a in authors) and nb_auth == AUTHOR, f'docx/pptx/pdf/notebook authors: {set(authors) | {nb_auth}}')

# ---------------------------------------------------------------- data/model integrity
check('CSV and model hashes match metadata', sha(ROOT / 'data/energy_efficiency.csv') == meta['data_sha256'] and sha(ROOT / 'models/energywise_linear_regression.joblib') == meta['model_sha256'],
      'data_sha256 and model_sha256 match')
check('Dataset size meets guideline (>=20 rows, >=2 features)', nbm['shape'][0] >= 20 and len(meta['features']) >= 2, f"{nbm['shape'][0]} rows, {len(meta['features'])} features")
check('At least 3 new predictions', len(nbm['new_predictions']) >= 3, f"{nbm['new_predictions']}")

# ---------------------------------------------------------------- console evidence + tests
im = Image.open(ROOT / 'reports/console_screenshot.png'); transcript = (ROOT / 'reports/console_transcript.txt').read_text(encoding='utf-8')
check('Real console screenshot + transcript', im.size[0] > 500 and 'Tahmini ısıtma yükü' in transcript and 'bir sayı değil' in transcript and 'Görüşmek üzere' in transcript,
      f'screenshot {im.size[0]}x{im.size[1]}; transcript shows prediction, invalid-input message and exit')
suite = ET.parse(ROOT / 'reports/test_results.xml').getroot()
suite = suite if suite.tag == 'testsuite' else suite.find('testsuite')
t, f, e = int(suite.get('tests')), int(suite.get('failures')), int(suite.get('errors'))
check('Automated tests', t >= 14 and f == 0 and e == 0, f'{t} tests, {f} failures, {e} errors')

# ---------------------------------------------------------------- zip
if '--zip' in sys.argv:
    if ZIP.is_file():
        with zipfile.ZipFile(ZIP) as z:
            bad = z.testzip(); names = z.namelist()
        leaked = [n for n in names if re.search(r'(^|/)(\.venv|__pycache__|\.pytest_cache|\.ipynb_checkpoints)(/|$)|(^|/)\.env$', n)]
        miss_zip = [f'EnergyWise/{r}' for r in REQUIRED if f'EnergyWise/{r}' not in names]
        check('ZIP integrity and contents', bad is None and not leaked and not miss_zip,
              f'{len(names)} files, {ZIP.stat().st_size / 1048576:.2f} MB, CRC ok={bad is None}, leaked={leaked or "none"}, missing={miss_zip or "none"}')
    else:
        check('ZIP integrity and contents', False, 'dist/EnergyWise_Teslim.zip not found')

ok = all(r[1] for r in results)
lines = ['EnergyWise - Delivery Verification', f'Generated: {time.strftime("%Y-%m-%d %H:%M:%S")}', '']
lines += [f'[{"PASS" if r[1] else "FAIL"}] {r[0]}: {r[2]}' for r in results]
lines += ['', f'OVERALL: {"PASS" if ok else "FAIL"} ({sum(r[1] for r in results)}/{len(results)})']
out = ROOT / 'reports' / ('zip_verification.txt' if '--zip' in sys.argv else 'delivery_verification.txt')
out.write_text('\n'.join(lines) + '\n', encoding='utf-8')
sys.stdout.reconfigure(encoding='utf-8'); print('\n'.join(lines))
sys.exit(0 if ok else 1)
