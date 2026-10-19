"""Build the 1-page report (DOCX -> PDF via Word) and the 5-slide deck (PPTX -> PDF via PowerPoint)
from the REAL outputs: reports/notebook_metrics.json, models/model_metadata.json, reports/figures/*.png,
reports/console_screenshot.png. No number is typed by hand.
Run: .venv/Scripts/python scripts/build_report.py   (then scripts/export_office.ps1 for the PDFs)
"""
import datetime as dt
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt as PPt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / 'reports' / 'figures'
AUTHOR = 'Tugay Sinan Ergün'
NB = json.loads((ROOT / 'reports' / 'notebook_metrics.json').read_text(encoding='utf-8'))
META = json.loads((ROOT / 'models' / 'model_metadata.json').read_text(encoding='utf-8'))
TR, TE, F = NB['train_metrics'], NB['test_metrics'], NB['findings']
BLUE, DARK, GREY = '2E86AB', '1B2A41', '5A6378'


def n(x, d=2):
    """Turkish decimal comma."""
    return f'{x:.{d}f}'.replace('.', ',')


# ------------------------------------------------------------------ texts shared by report and deck
FINDINGS = [
    f"Yükseklik en belirleyici özellik: 3,5 m binalarda ortalama ısıtma yükü {n(F['yukseklik_ortalama_3_5'])}, 7,0 m binalarda {n(F['yukseklik_ortalama_7_0'])}.",
    f"Cam alanı arttıkça yük artıyor: cam yokken ortalama {n(F['cam_0_ortalama'])}, 0,40 cam oranında {n(F['cam_0_40_ortalama'])}.",
    f"Yönün etkisi çok küçük: yönler arasındaki en büyük ortalama fark {n(F['yon_ortalama_fark'], 3)}.",
    f"En güçlü pozitif korelasyon {F['en_guclu_pozitif_korelasyon'][0]} (r = {n(F['en_guclu_pozitif_korelasyon'][1], 3)}), "
    f"en güçlü negatif korelasyon {F['en_guclu_negatif_korelasyon'][0]} (r = {n(F['en_guclu_negatif_korelasyon'][1], 3)}).",
]
METRIC_ROWS = [('MAE', n(TR['MAE'], 3), n(TE['MAE'], 3)), ('RMSE', n(TR['RMSE'], 3), n(TE['RMSE'], 3)), ('R²', n(TR['R2'], 4), n(TE['R2'], 4))]


# ------------------------------------------------------------------ 1-page report
def report():
    doc = Document()
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.6); s.top_margin = s.bottom_margin = Cm(1.3)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); st.paragraph_format.space_after = Pt(2)

    def head(text, size=10.5, color=BLUE):
        p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(1)

    def para(text):
        p = doc.add_paragraph()
        for i, part in enumerate(text.split('**')):
            r = p.add_run(part); r.bold = i % 2 == 1
        return p

    t = doc.add_paragraph(); r = t.add_run('EnergyWise — Bina Isıtma Yükü Tahmini'); r.bold = True; r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(DARK); t.paragraph_format.space_after = Pt(0)
    para(f'Proje raporu · Yapay Zeka İçin Basitleştirilmiş Python · {AUTHOR} (GitHub: TCoderr) · 19 Ekim 2026')

    head('1. Problem')
    para('Bir binanın mimari özelliklerinden **ısıtma yükünü** (heating_load) tahmin etmek. Problem türü: **regresyon**. '
         'Tasarım aşamasında farklı biçim ve cam seçeneklerini hızlıca karşılaştırmak için ayrıntılı simülasyona göre daha hızlı bir ön tahmin sağlar.')
    head('2. Veri')
    para(f"UCI Machine Learning Repository — Energy Efficiency (Tsanas & Xifara, 2012; CC BY 4.0; doi:10.24432/C51307). "
         f"{NB['shape'][0]} satır, {NB['shape'][1]} sütun; 8 girdi, hedef: heating_load; eksik değer: {NB['missing_total']}. "
         '**Veri gerçek binalardan ölçülmemiştir**; Ecotect\'te simüle edilen 12 bina şeklinden üretilmiştir. '
         'İkinci hedef cooling_load, hedef sızıntısını önlemek için girdi olarak kullanılmadı.')
    head('3. Keşif bulguları')
    for f in FINDINGS:
        doc.add_paragraph(f, style='List Bullet').paragraph_format.space_after = Pt(0)
    head('4. Model')
    para(f"sklearn LinearRegression; train_test_split test_size=0,2, random_state=42 → eğitim {NB['n_train']}, test {NB['n_test']} satır. "
         'Doğrusal regresyon ölçekten etkilenmediği için ölçeklendirme yapılmadı; dönüşüm olmadığından test verisinden bilgi sızmaz.')
    tbl = doc.add_table(rows=4, cols=3); tbl.style = 'Light Grid Accent 1'
    for i, row in enumerate([('Metrik', 'Eğitim', 'Test')] + METRIC_ROWS):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = v
            for p in tbl.cell(i, j).paragraphs:
                for r in p.runs: r.font.size = Pt(8.5)
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(FIG / '05_gercek_tahmin.png'), height=Cm(5.2))
    pic.add_run().add_picture(str(FIG / '02_kompaktlik_isitma.png'), height=Cm(5.2))
    head('5. Sonuç ve çıkarım')
    p3 = ', '.join(n(v) for v in NB['new_predictions'])
    para(f"Test R² = {n(TE['R2'], 4)} ve MAE = {n(TE['MAE'], 3)}: model ısıtma yükündeki değişimin büyük kısmını açıklıyor. "
         f"Eğitim ve test sonuçları birbirine yakın, belirgin aşırı öğrenme yok. Veri setinde bulunmayan 3 yeni bina için tahminler: {p3}. "
         'tahmin_et() fonksiyonu terminalde girdileri doğrulayarak tahmin üretiyor. '
         '**Sınırlar:** Model doğrusal olduğu için yükseklik ve kompaktlığın doğrusal olmayan etkisini tam yakalayamıyor. '
         'Veri 12 simüle bina şekliyle sınırlı ve gerçek binalara genellenemez. '
         '**İyileştirme önerileri:** kodlu değişkenleri one-hot kodlamak, etkileşim terimleri eklemek, karar ağacı veya rastgele orman denemek.')
    cp = doc.core_properties; cp.author = cp.last_modified_by = AUTHOR; cp.title = 'EnergyWise Proje Raporu'
    cp.created = cp.modified = dt.datetime(2026, 10, 19, 12, 0)
    out = ROOT / 'reports' / 'EnergyWise_Proje_Raporu.docx'; doc.save(out); return out


# ------------------------------------------------------------------ 5-slide deck
def deck():
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    rgb = PRGB.from_string

    def slide(title):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)); bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(BLUE); bar.line.fill.background()
        text(s, title, 0.6, 0.3, 12, 0.8, 28, DARK, True)
        text(s, f'EnergyWise · {AUTHOR}', 0.6, 7.0, 6, 0.3, 10, GREY)
        return s

    def text(s, lines, x, y, w, h, size, color=DARK, bold=False):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = PPt(8)
            r = p.add_run(); r.text = line; r.font.size = PPt(size); r.font.bold = bold; r.font.name = 'Arial'
            r.font.color.rgb = rgb(color)
        return tb

    s = slide('Problem ve veri')
    text(s, ['Problem: Binanın mimari özelliklerinden ısıtma yükünü tahmin etmek (regresyon).',
             f"Veri: UCI Energy Efficiency — {NB['shape'][0]} satır, 8 girdi, hedef heating_load, eksik değer {NB['missing_total']}.",
             'Lisans: CC BY 4.0 · Tsanas & Xifara (2012).',
             'Veri simülasyondur: Ecotect\'te 12 bina şekli; gerçek bina ölçümü değildir.',
             'cooling_load ikinci hedef olduğu için girdi olarak kullanılmadı.'], 0.6, 1.4, 6.3, 5, 17)
    s.shapes.add_picture(str(FIG / '01_hedef_dagilimi.png'), Inches(7.1), Inches(1.5), width=Inches(5.8))

    s = slide('Keşif bulguları')
    text(s, FINDINGS, 0.6, 1.4, 6.2, 5.4, 15)
    s.shapes.add_picture(str(FIG / '02_kompaktlik_isitma.png'), Inches(7.0), Inches(1.4), width=Inches(5.9))

    s = slide('Model: Linear Regression')
    text(s, [f"Eğitim {NB['n_train']} / test {NB['n_test']} satır (test_size=0,2, random_state=42)."] +
         [f'{m}: eğitim {a} · test {b}' for m, a, b in METRIC_ROWS] +
         ['Eğitim ve test sonuçları yakın: belirgin aşırı öğrenme yok.'], 0.6, 1.4, 6.0, 5.4, 17)
    s.shapes.add_picture(str(FIG / '05_gercek_tahmin.png'), Inches(7.0), Inches(1.3), height=Inches(5.3))  # 5.83 in wide

    s = slide('Yeni tahminler ve konsol uygulaması')
    text(s, ['Veri setinde bulunmayan 3 yeni bina (cam oranı 0,05 / 0,15 / 0,30):'] +
         [f'Örnek {i + 1}: tahmini ısıtma yükü {n(v)}' for i, v in enumerate(NB['new_predictions'])] +
         ['tahmin_et(): input() ile girdi alır, doğrular, f-string ile yazdırır; q ile çıkış.'], 0.6, 1.4, 5.2, 5.4, 16)
    shot = ROOT / 'reports' / 'console_screenshot.png'
    if shot.exists():
        s.shapes.add_picture(str(shot), Inches(7.2), Inches(1.3), height=Inches(5.5))  # keep inside the 7.5 in slide height

    s = slide('Sınırlar ve sonuç')
    text(s, [f"Test R² {n(TE['R2'], 4)}, MAE {n(TE['MAE'], 3)}: hızlı bir ön tahmin için yeterli.",
             'Doğrusal model; yükseklik ve kompaktlığın doğrusal olmayan etkisini tam yakalayamıyor.',
             'Veri 12 simüle bina şekliyle sınırlı; gerçek binalara ve farklı iklimlere genellenemez.',
             'İyileştirme: one-hot kodlama, etkileşim terimleri, ağaç tabanlı modeller, çapraz doğrulama.'], 0.6, 1.4, 6.2, 5.4, 17)
    s.shapes.add_picture(str(FIG / '06_hata_dagilimi.png'), Inches(7.0), Inches(1.5), width=Inches(5.9))

    cp = prs.core_properties; cp.author = cp.last_modified_by = AUTHOR; cp.title = 'EnergyWise Sunum'
    cp.created = cp.modified = dt.datetime(2026, 10, 19, 12, 0)
    out = ROOT / 'presentation' / 'EnergyWise_Sunum.pptx'; prs.save(out); return out


if __name__ == '__main__':
    print(report()); print(deck())
